"""Génère le script oral et la fiche de révision en PDF et DOCX à partir du Markdown.

    python tools/build_documents.py

Sources : docs/Script_oral_5min.md et docs/Fiche_revision_ultime.md.
La police DejaVu Sans est celle fournie avec matplotlib : même rendu en local
et dans GitHub Actions, avec tous les caractères (→, ≤, «, é…).
"""

from __future__ import annotations

from html import escape
from pathlib import Path
import re

import matplotlib
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
DOCUMENTS = [
    ("Script_oral_5min.md", "Script_oral_5min"),
    ("Fiche_revision_ultime.md", "Fiche_revision_ultime"),
]

INK = "#0B1F2A"
ORANGE = "#EB6834"
MUTED = "#5B636A"
LIGHT = "#F3F5F7"
LINE = "#D9DDE1"


# ---------------------------------------------------------------------------
# Markdown minimal : titres, paragraphes, citations, listes, tableaux, séparateurs
# ---------------------------------------------------------------------------


def blocks(path: Path):
    lines = path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        raw = lines[i].rstrip()
        stripped = raw.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            yield ("table", rows)
            continue
        if stripped.startswith(">"):
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip()[1:].strip())
                i += 1
            yield ("quote", quote)
            continue
        if stripped == "---":
            yield ("rule", None)
        elif stripped.startswith("# "):
            yield ("h1", stripped[2:])
        elif stripped.startswith("## "):
            yield ("h2", stripped[3:])
        elif stripped.startswith("### "):
            yield ("h3", stripped[4:])
        elif re.match(r"^\d+\.\s", stripped):
            yield ("number", re.sub(r"^\d+\.\s+", "", stripped), int(stripped.split(".")[0]))
        elif stripped.startswith("- "):
            yield ("bullet", stripped[2:])
        else:
            yield ("text", stripped)
        i += 1


INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def inline_parts(text: str):
    """Découpe en (texte, gras, code)."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            yield part[2:-2], True, False
        elif part.startswith("`"):
            yield part[1:-1], False, True
        else:
            yield part, False, False


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#"))


def _shade(cell, hex_color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_color.lstrip("#"))
    props.append(shading)


def _runs(paragraph, text, size=None, color=None, bold=False, italic=False):
    for chunk, is_bold, is_code in inline_parts(text):
        run = paragraph.add_run(chunk)
        run.bold = bold or is_bold
        run.italic = italic
        if is_code:
            run.font.name = "Consolas"
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = _rgb(color)


def build_docx(source: Path, target: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, side, Mm(18))
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for name, size in (("Heading 1", 16), ("Heading 2", 13)):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = _rgb(INK)

    for block in blocks(source):
        kind = block[0]
        if kind == "h1":
            p = doc.add_paragraph()
            _runs(p, block[1], size=22, color=INK, bold=True)
        elif kind == "h2":
            doc.add_heading(block[1], level=1)
        elif kind == "h3":
            doc.add_heading(block[1], level=2)
        elif kind == "bullet":
            _runs(doc.add_paragraph(style="List Bullet"), block[1])
        elif kind == "number":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Mm(6)
            p.paragraph_format.first_line_indent = Mm(-6)
            _runs(p, f"{block[2]}.  {block[1]}")
        elif kind == "quote":
            table = doc.add_table(rows=1, cols=1)
            cell = table.cell(0, 0)
            _shade(cell, LIGHT)
            cell.paragraphs[0].text = ""
            for i, line in enumerate(block[1]):
                p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
                _runs(p, line, size=11.5, color=INK)
            doc.add_paragraph()
        elif kind == "table":
            rows = block[1]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    cell = table.cell(r, c)
                    cell.paragraphs[0].text = ""
                    _runs(cell.paragraphs[0], value, size=10, color="#FFFFFF" if r == 0 else None, bold=r == 0)
                    if r == 0:
                        _shade(cell, INK)
            doc.add_paragraph()
        elif kind == "rule":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _runs(p, "·  ·  ·", color=MUTED)
        elif kind == "text":
            _runs(doc.add_paragraph(), block[1])

    doc.core_properties.author = "Romain, Lisa, Yannis"
    doc.core_properties.title = source.stem.replace("_", " ")
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def register_fonts() -> tuple[str, str, str]:
    font_dir = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
    pdfmetrics.registerFont(TTFont("DejaVu", str(font_dir / "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(font_dir / "DejaVuSans-Bold.ttf")))
    pdfmetrics.registerFont(TTFont("DejaVuMono", str(font_dir / "DejaVuSansMono.ttf")))
    pdfmetrics.registerFontFamily("DejaVu", normal="DejaVu", bold="DejaVu-Bold", italic="DejaVu", boldItalic="DejaVu-Bold")
    return "DejaVu", "DejaVu-Bold", "DejaVuMono"


def _markup(text: str, mono: str) -> str:
    out = []
    for chunk, bold, code in inline_parts(text):
        chunk = escape(chunk)
        if bold:
            chunk = f"<b>{chunk}</b>"
        if code:
            chunk = f'<font face="{mono}">{chunk}</font>'
        out.append(chunk)
    return "".join(out)


def build_pdf(source: Path, target: Path) -> None:
    regular, bold, mono = register_fonts()
    body = ParagraphStyle("body", fontName=regular, fontSize=9.6, leading=13.6, spaceAfter=4, textColor=colors.HexColor(INK), alignment=TA_LEFT)
    small = ParagraphStyle("small", parent=body, fontSize=8.6, leading=11.4, spaceAfter=0)
    h1 = ParagraphStyle("h1", parent=body, fontName=bold, fontSize=19, leading=24, spaceAfter=6)
    h2 = ParagraphStyle("h2", parent=body, fontName=bold, fontSize=13, leading=17, spaceBefore=10, spaceAfter=5, textColor=colors.HexColor(ORANGE))
    h3 = ParagraphStyle("h3", parent=body, fontName=bold, fontSize=11, leading=14, spaceBefore=6, spaceAfter=3)
    bullet = ParagraphStyle("bullet", parent=body, leftIndent=12, bulletIndent=2, spaceAfter=2)
    quote = ParagraphStyle("quote", parent=body, fontSize=10.2, leading=14.5, spaceAfter=2)

    story = []
    for block in blocks(source):
        kind = block[0]
        if kind == "h1":
            story.append(Paragraph(_markup(block[1], mono), h1))
        elif kind == "h2":
            story.append(Paragraph(_markup(block[1], mono), h2))
        elif kind == "h3":
            story.append(Paragraph(_markup(block[1], mono), h3))
        elif kind == "bullet":
            story.append(Paragraph(_markup(block[1], mono), bullet, bulletText="•"))
        elif kind == "number":
            story.append(Paragraph(_markup(block[1], mono), bullet, bulletText=f"{block[2]}."))
        elif kind == "quote":
            inner = [[Paragraph(_markup(line, mono), quote)] for line in block[1]]
            box = Table(inner, colWidths=[176 * mm])
            box.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(LIGHT)),
                ("LEFTPADDING", (0, 0), (-1, -1), 9), ("RIGHTPADDING", (0, 0), (-1, -1), 9),
                ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, 0), 8), ("BOTTOMPADDING", (0, -1), (-1, -1), 8),
            ]))
            story += [box, Spacer(1, 6)]
        elif kind == "table":
            rows = block[1]
            header = ParagraphStyle("th", parent=small, fontName=bold, textColor=colors.white)
            data = [[Paragraph(_markup(c, mono), header if r == 0 else small) for c in row] for r, row in enumerate(rows)]
            table = Table(data, repeatRows=1, hAlign="LEFT", colWidths=_col_widths(rows))
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(INK)),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(LIGHT)]),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor(LINE)),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            story += [KeepTogether(table), Spacer(1, 7)]
        elif kind == "rule":
            story.append(Spacer(1, 8))
        elif kind == "text":
            story.append(Paragraph(_markup(block[1], mono), body))

    def page(canvas, doc):
        canvas.saveState()
        canvas.setFont(regular, 7.5)
        canvas.setFillColor(colors.HexColor(MUTED))
        canvas.drawString(17 * mm, 10 * mm, "Une vie de fourmi · Romain · Lisa · Yannis")
        canvas.drawRightString(A4[0] - 17 * mm, 10 * mm, str(doc.page))
        canvas.restoreState()

    target.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(
        str(target), pagesize=A4, leftMargin=17 * mm, rightMargin=17 * mm, topMargin=15 * mm, bottomMargin=17 * mm,
        title=source.stem.replace("_", " "), author="Romain, Lisa, Yannis",
    ).build(story, onFirstPage=page, onLaterPages=page)


def _col_widths(rows):
    """Largeurs proportionnelles au contenu, sur 176 mm."""
    cols = len(rows[0])
    lengths = [max(len(re.sub(r"[`*]", "", row[c])) for row in rows) for c in range(cols)]
    lengths = [min(max(l, 8), 60) for l in lengths]
    total = sum(lengths)
    return [176 * mm * l / total for l in lengths]


def main() -> None:
    for source_name, stem in DOCUMENTS:
        source = DOCS / source_name
        build_docx(source, DOCS / f"{stem}.docx")
        build_pdf(source, DOCS / f"{stem}.pdf")
        print(f"{stem}.pdf / .docx générés")


if __name__ == "__main__":
    main()
